from pathlib import Path
import csv
from openpyxl import load_workbook

def load_keywords(file_path: str | Path) -> list[str]:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _load_keywords_from_csv(path)
    if suffix == ".xlsx":
        return _load_keywords_from_xlsx(path)
    raise ValueError("Only .csv and .xlsx files are supported.")


def _dedupe_keywords(values: list[str]) -> list[str]:
    keywords: list[str] = []
    seen: set[str] = set()
    for value in values:
        keyword = value.strip()
        if not keyword or keyword in seen:
            continue
        seen.add(keyword)
        keywords.append(keyword)
    if not keywords:
        raise ValueError("No non-empty keyword rows were found.")
    return keywords


def _load_keywords_from_csv(path: Path) -> list[str]:
    decode_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8-sig", "cp949"):
        try:
            with path.open("r", encoding=encoding, newline="") as file:
                reader = csv.DictReader(file)
                fieldnames = reader.fieldnames or []
                keyword_header = fieldnames[0] if fieldnames else ""
                if not keyword_header or keyword_header.strip().lower() != "keyword":
                    raise ValueError("The first CSV header must be keyword.")
                return _dedupe_keywords([(row.get(keyword_header) or "") for row in reader])
        except UnicodeDecodeError as exc:
            decode_error = exc
    raise ValueError("CSV encoding must be UTF-8 or CP949.") from decode_error


def _load_keywords_from_xlsx(path: Path) -> list[str]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required for XLSX files.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValueError("XLSX file is empty.")
    headers = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    try:
        keyword_index = headers.index("keyword")
    except ValueError as exc:
        raise ValueError("XLSX header must include keyword.") from exc
    return _dedupe_keywords(
        [
            str(row[keyword_index]) if keyword_index < len(row) and row[keyword_index] is not None else ""
            for row in rows[1:]
        ]
    )
